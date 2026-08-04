"""
ISS Substituto worker — 4 parallel Playwright browsers.
Each worker gets its own browser instance and processes a chunk of companies.
Logs from all workers feed into a shared queue for SSE streaming.
Checkpoint writes are thread-safe via a lock.
"""
import base64
import hashlib
import io
import json
import re
import threading
import time
import random
from datetime import datetime
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed

from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE        = "https://tubarao-sc.prefeituramoderna.com.br/meuiss_new"
MAX_RETRIES = 3
MAX_WORKERS = 4
CHECKPOINT  = Path("pm_checkpoint.json")

checkpoint_lock = threading.Lock()


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def parse_money(text: str) -> float:
    text = text.strip().replace(".", "").replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def rand_delay(lo=1.0, hi=2.5):
    return random.uniform(lo, hi)


# ---------------------------------------------------------------------------
# HTML parsing
# ---------------------------------------------------------------------------

def is_period_open(html: str) -> bool:
    return "Período em Aberto" in html or "Periodo em Aberto" in html


def get_period_breakdown(html: str) -> dict:
    soup   = BeautifulSoup(html, "html.parser")
    totals = {k: 0.0 for k in ("Normal", "Cancelado", "Retido", "Substituto", "Tomador")}
    for tbody in soup.find_all("tbody"):
        rows = tbody.find_all("tr")
        target = None
        for row in rows:
            tds = row.find_all("td")
            if len(tds) == 11 and any(re.match(r"[\d.]+,\d+", td.get_text(strip=True)) for td in tds):
                target = tbody
                break
        if target:
            for row in target.find_all("tr"):
                tds = row.find_all("td")
                if len(tds) != 11:
                    continue
                totals["Normal"]     += parse_money(tds[2].get_text())
                totals["Cancelado"]  += parse_money(tds[4].get_text())
                totals["Retido"]     += parse_money(tds[6].get_text())
                totals["Substituto"] += parse_money(tds[8].get_text())
                totals["Tomador"]    += parse_money(tds[10].get_text())
            break
    return totals


# ---------------------------------------------------------------------------
# Browser / portal helpers
# ---------------------------------------------------------------------------

def create_browser(visible=False):
    from playwright.sync_api import sync_playwright
    p       = sync_playwright().start()
    browser = p.chromium.launch(headless=not visible, slow_mo=80)
    context = browser.new_context(
        user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                   "(KHTML, like Gecko) Chrome/124.0 Safari/537.36",
        viewport={"width": 1280, "height": 720},
    )
    return p, browser, context


def do_login(context, user: str, password: str):
    senha_hash = hashlib.md5(password.encode("utf-8")).hexdigest()
    page = context.new_page()
    page.goto(f"{BASE}/index.php?out=2", wait_until="networkidle", timeout=60000)
    page.wait_for_timeout(3000)
    if page.locator("#login_iss").count() == 0:
        page.wait_for_timeout(5000)
    if page.locator("#login_iss").count() == 0:
        raise RuntimeError("Login form not found — are you on a Brazilian IP?")
    page.fill("#login_iss", user)
    page.fill("#senha_iss_digite", password)
    page.evaluate(
        "(hash) => { document.querySelector('input[name=\"senha_iss\"]').value = hash; }",
        senha_hash,
    )
    page.click("input[type=submit][value='Acessar Sistema']")
    page.wait_for_load_state("networkidle", timeout=60000)
    page.wait_for_timeout(2000)
    if page.locator("#login_iss").is_visible():
        raise RuntimeError("Login failed — wrong credentials or Cloudflare block.")
    return page


def login(user: str, password: str, visible: bool = False):
    """Used by nfse_worker and guias_worker."""
    p, browser, context = create_browser(visible)
    page = do_login(context, user, password)
    return p, browser, context, page


def get_company_list(page) -> list[dict]:
    page.evaluate("() => { mostrarConteudo('iss-clientes_contador.php'); }")
    page.wait_for_timeout(2000)
    page.wait_for_selector("select#clientes", timeout=10000)
    return page.eval_on_selector_all(
        "select#clientes option",
        "opts => opts.map(o => ({id: o.value, name: o.text})).filter(o => o.id && o.id !== '0')"
    )


def go_to_empresas(page):
    try:
        page.evaluate("() => { document.form_painel.submit(); }")
        page.wait_for_load_state("networkidle", timeout=15000)
        page.wait_for_timeout(1000)
    except Exception:
        pass
    page.evaluate("() => { mostrarConteudo('iss-clientes_contador.php'); }")
    page.wait_for_timeout(2000)


def get_company_data(page, company_id: str, target_period: str | None) -> dict:
    go_to_empresas(page)
    page.wait_for_selector("select#clientes", timeout=10000)
    page.select_option("select#clientes", company_id)
    page.wait_for_timeout(800)
    page.click("input[value='Acessar Movimento Econômico']")
    page.wait_for_load_state("networkidle", timeout=30000)
    page.wait_for_timeout(1200)
    page.evaluate("() => { mostrarConteudo('iss-consulta_periodos.php'); }")
    page.wait_for_timeout(2000)
    page.wait_for_selector("select#id_movimento", timeout=10000)

    period_options = page.eval_on_selector_all(
        "select#id_movimento option",
        "opts => opts.map(o => o.value).filter(v => v && v !== '0')"
    )
    if not period_options:
        return {"mesano": None, "breakdown": None}

    if target_period:
        if target_period not in period_options:
            return {"mesano": target_period, "breakdown": None, "skipped": "not_found"}
        page.select_option("select#id_movimento", target_period)
        page.wait_for_timeout(2500)
        html = page.content()
        if is_period_open(html):
            return {"mesano": target_period, "breakdown": None, "skipped": "open"}
        return {"mesano": target_period, "breakdown": get_period_breakdown(html)}
    else:
        for mesano in period_options:
            page.select_option("select#id_movimento", mesano)
            page.wait_for_timeout(2500)
            html = page.content()
            if not is_period_open(html):
                return {"mesano": mesano, "breakdown": get_period_breakdown(html)}
        return {"mesano": period_options[0], "breakdown": {k: 0.0 for k in ("Normal","Cancelado","Retido","Substituto","Tomador")}}


# ---------------------------------------------------------------------------
# Checkpoint (thread-safe)
# ---------------------------------------------------------------------------

def load_checkpoint() -> dict:
    if CHECKPOINT.exists():
        return json.loads(CHECKPOINT.read_text(encoding="utf-8"))
    return {"done": {}}


def save_checkpoint(state: dict):
    with checkpoint_lock:
        CHECKPOINT.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")


# ---------------------------------------------------------------------------
# Excel report
# ---------------------------------------------------------------------------

def build_report(state: dict, target_period: str | None) -> bytes:
    rows_yes, rows_no = [], []
    for cid, data in state["done"].items():
        total = data.get("substituto_total")
        if data.get("error") or data.get("skipped") or total is None:
            continue
        m          = data.get("mesano", "")
        period_str = f"{m[:2]}/{m[2:]}" if m and len(m) == 6 else ""
        row = {
            "name":     data["name"],
            "cmc":      cid,
            "has_sub":  "Sim" if total > 0 else "Não",
            "sub":      total,
            "has_nor":  "Sim" if data.get("normal_total", 0) > 0 else "Não",
            "nor":      data.get("normal_total", 0.0),
            "period":   period_str,
        }
        (rows_yes if total > 0 else rows_no).append(row)

    rows_yes.sort(key=lambda r: r["sub"], reverse=True)
    rows_no.sort(key=lambda r: r["name"])
    all_rows = rows_yes + rows_no

    wb  = Workbook()
    ws  = wb.active
    ws.title = "ISS Substituto"

    hf    = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    hfill = PatternFill("solid", start_color="C0392B")
    tf    = Font(name="Arial", bold=True, size=11, color="C0392B")
    nf    = Font(name="Arial", size=10)
    bf    = Font(name="Arial", bold=True, size=10)
    yf    = PatternFill("solid", start_color="FDEBD0")
    af    = PatternFill("solid", start_color="F8F9FA")
    wf    = PatternFill("solid", start_color="FFFFFF")
    totf  = PatternFill("solid", start_color="F9E79F")
    ctr   = Alignment(horizontal="center", vertical="center")
    lft   = Alignment(horizontal="left",   vertical="center")
    rgt   = Alignment(horizontal="right",  vertical="center")
    thin  = Side(style="thin", color="BFBFBF")
    bdr   = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = "#,##0.00"

    period_label = f" — Período {target_period[:2]}/{target_period[2:]}" if target_period else ""
    ws.merge_cells("A1:G1")
    ws["A1"] = f"ISS Substituto — Relatório de Presença{period_label}"
    ws["A1"].font = tf; ws["A1"].alignment = lft
    ws.merge_cells("A2:G2")
    ws["A2"] = (f"Gerado em {datetime.now().strftime('%d/%m/%Y %H:%M')}  |  "
                f"{len(all_rows)} empresas  |  {len(rows_yes)} com ISS Substituto")
    ws["A2"].font = Font(name="Arial", size=9, color="7F8C8D"); ws["A2"].alignment = lft

    for col, h in enumerate(["Empresa","CMC","ISS Substituto","Valor Substituto","ISS Normal","Valor Normal","Período"], 1):
        c = ws.cell(row=3, column=col, value=h)
        c.font = hf; c.fill = hfill; c.alignment = ctr; c.border = bdr
    ws.row_dimensions[3].height = 16

    for i, r in enumerate(all_rows, start=4):
        fill = yf if r["has_sub"] == "Sim" else (af if i % 2 == 0 else wf)
        for col, val in enumerate([r["name"],r["cmc"],r["has_sub"],r["sub"],r["has_nor"],r["nor"],r["period"]], 1):
            c = ws.cell(row=i, column=col, value=val)
            c.font = nf; c.fill = fill; c.border = bdr
            if col in (4, 6): c.number_format = money; c.alignment = rgt
            elif col == 1:    c.alignment = lft
            else:             c.alignment = ctr

    for letter, w in {"A":55,"B":12,"C":16,"D":20,"E":14,"F":18,"G":12}.items():
        ws.column_dimensions[letter].width = w
    ws.freeze_panes = "A4"

    # Skipped / errors sheet
    skipped = [(cid, d) for cid, d in state["done"].items() if d.get("skipped") or d.get("error")]
    if skipped:
        ws2 = wb.create_sheet("Avisos")
        ws2.append(["Empresa", "CMC", "Situação"])
        ws2.column_dimensions["A"].width = 50
        ws2.column_dimensions["C"].width = 30
        for cid, d in skipped:
            ws2.append([d.get("name",""), cid, d.get("error") or d.get("skipped","")])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Per-worker function (runs in its own thread with its own browser)
# ---------------------------------------------------------------------------

def worker_run(worker_id: int, companies: list[dict], user: str, password: str,
               target_period: str | None, state: dict, log_queue, visible: bool = False):
    """
    Processes a chunk of companies in its own browser.
    Pushes log strings to log_queue.
    Updates state dict directly (thread-safe via checkpoint_lock for writes).
    """
    def log(msg):
        line = f"[{datetime.now().strftime('%H:%M:%S')}] [W{worker_id}] {msg}"
        log_queue.put(line)

    period_label = f"{target_period[:2]}/{target_period[2:]}" if target_period else "mais recente"

    try:
        p, browser, context = create_browser(visible)
        page = do_login(context, user, password)
        log(f"Login OK — {len(companies)} empresas")
    except Exception as e:
        log(f"ERRO login: {e}")
        return

    try:
        for i, comp in enumerate(companies, 1):
            cid  = comp["id"]
            name = comp["name"]

            with checkpoint_lock:
                if cid in state["done"]:
                    log(f"{name} — checkpoint (skip)")
                    continue

            log(f"[{i}/{len(companies)}] {name}")

            for attempt in range(1, MAX_RETRIES + 1):
                try:
                    result  = get_company_data(page, cid, target_period)
                    skipped = result.get("skipped")

                    if skipped == "not_found":
                        entry = {"name": name, "substituto_total": 0.0, "normal_total": 0.0,
                                 "mesano": target_period, "skipped": "not_found"}
                        log(f"  -> Período {period_label} não disponível")
                    elif skipped == "open":
                        entry = {"name": name, "substituto_total": 0.0, "normal_total": 0.0,
                                 "mesano": target_period, "skipped": "open"}
                        log(f"  -> Período em aberto")
                    elif result["breakdown"] is None:
                        entry = {"name": name, "substituto_total": 0.0, "normal_total": 0.0, "mesano": None}
                        log(f"  -> Sem períodos")
                    else:
                        bd  = result["breakdown"]
                        sub = bd["Substituto"]
                        nor = bd["Normal"]
                        entry = {"name": name, "substituto_total": sub, "normal_total": nor,
                                 "mesano": result["mesano"]}
                        if sub > 0:
                            log(f"  -> ISS Substituto: R$ {sub:,.2f} ✓")
                        else:
                            log(f"  -> Sem ISS Substituto")

                    with checkpoint_lock:
                        state["done"][cid] = entry

                    break

                except Exception as exc:
                    log(f"  -> [Tentativa {attempt}] {exc}")
                    if attempt == MAX_RETRIES:
                        with checkpoint_lock:
                            state["done"][cid] = {"name": name, "substituto_total": None,
                                                   "normal_total": None, "mesano": None, "error": str(exc)}
                    else:
                        time.sleep(2)

            # Checkpoint every 10 companies per worker
            if i % 10 == 0:
                save_checkpoint(state)

            if i < len(companies):
                time.sleep(rand_delay(0.8, 1.8))

    except Exception as e:
        log(f"ERRO fatal no worker: {e}")
    finally:
        try:
            browser.close()
            p.stop()
        except Exception:
            pass
        save_checkpoint(state)
        log("Worker encerrado")


# ---------------------------------------------------------------------------
# Main generator
# ---------------------------------------------------------------------------

def run_iss_report(user: str, password: str, target_period: str | None,
                   limit: int | None = None, resume: bool = False,
                   workers: int = MAX_WORKERS, visible: bool = False):
    """
    Generator. Yields log strings for SSE streaming.
    Final yields: __DONE__:code, __XLSX__:base64
    """
    import queue as q_module

    def log(msg):
        return f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"

    yield log(f"Iniciando {workers} workers...")

    # Fetch company list with a temporary browser
    try:
        p0, b0, ctx0 = create_browser(visible)
        page0 = do_login(ctx0, user, password)
        companies = get_company_list(page0)
        b0.close(); p0.stop()
    except Exception as e:
        yield log(f"ERRO — Login/lista: {e}")
        yield "__DONE__:1"
        return

    yield log(f"{len(companies)} empresas encontradas.")

    if limit:
        companies = companies[:limit]
        yield log(f"Limitado a {limit} empresas.")

    period_label = f"{target_period[:2]}/{target_period[2:]}" if target_period else "mais recente"
    yield log(f"Período alvo: {period_label} | {workers} workers paralelos")

    state = load_checkpoint() if resume else {"done": {}}
    if resume:
        already = sum(1 for cid in state["done"] if any(c["id"] == cid for c in companies))
        yield log(f"Resumindo — {already} empresas já no checkpoint.")

    # Split companies into chunks for each worker
    pending = [c for c in companies if c["id"] not in state["done"]]
    yield log(f"{len(pending)} empresas para processar ({len(companies)-len(pending)} já no checkpoint)")

    # Distribute evenly across workers
    chunks = [[] for _ in range(workers)]
    for i, comp in enumerate(pending):
        chunks[i % workers].append(comp)

    log_queue = q_module.Queue()
    futures   = []

    with ThreadPoolExecutor(max_workers=workers) as executor:
        for wid, chunk in enumerate(chunks, 1):
            if not chunk:
                continue
            futures.append(executor.submit(
                worker_run, wid, chunk, user, password,
                target_period, state, log_queue, visible
            ))

        # Stream logs while workers are running
        active = len(futures)
        while active > 0:
            # Drain the queue
            while True:
                try:
                    msg = log_queue.get_nowait()
                    yield msg
                except q_module.Empty:
                    break

            # Check if any futures are done
            still_running = sum(1 for f in futures if not f.done())
            if still_running < active:
                active = still_running

            if active > 0:
                time.sleep(0.3)

        # Final drain
        while True:
            try:
                yield log_queue.get_nowait()
            except q_module.Empty:
                break

    save_checkpoint(state)
    substituto_count = sum(1 for d in state["done"].values()
                           if (d.get("substituto_total") or 0) > 0)
    yield log(f"Concluído — {substituto_count} empresas com ISS Substituto no período {period_label}.")
    yield "__DONE__:0"

    xlsx_bytes = build_report(state, target_period)
    yield f"__XLSX__:{base64.b64encode(xlsx_bytes).decode()}"
