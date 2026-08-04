"""
NFS-e PM worker — downloads movimento_economico XML per company/period from
Prefeitura Moderna portal, parses retention data, generates folder structure + XLSX.

Flow (confirmed from HAR):
  1. Select company → Acessar Movimento Econômico  (same as ISS)
  2. mostrarConteudo('iss-exportar_dados.php')
  3. POST gerar_xml.php with anomes_ini/fim + lancamentos=T
  4. Response is direct XML file download (application/octet-stream)

Period format: YYYYM for months 1-9, YYYYMM for 10-12
  e.g. Jun 2026 → "20266", Dec 2025 → "202512"
"""
import io
import re
import time
import random
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

BASE = "https://tubarao-sc.prefeituramoderna.com.br/meuiss_new"


# ---------------------------------------------------------------------------
# Period format conversion
# MMYYYY (UI input) → YYYYM portal format
# ---------------------------------------------------------------------------

def to_portal_period(mmyyyy: str) -> str:
    """Convert MMYYYY → portal's YYYYM format.
    e.g. '062026' → '20266', '122025' → '202512'
    """
    mm   = mmyyyy[:2].lstrip("0") or "0"
    yyyy = mmyyyy[2:]
    return f"{yyyy}{mm}"


# ---------------------------------------------------------------------------
# Money / safe name helpers
# ---------------------------------------------------------------------------

def fm(val) -> float:
    try:
        return float(str(val or "0").replace(".", "").replace(",", "."))
    except ValueError:
        return 0.0


def has_federal_retencao(n: dict) -> bool:
    return (n["v_irrf"] > 0 or n["v_csll"] > 0 or
            n["v_pis"]  > 0 or n["v_cofins"] > 0 or n["v_inss"] > 0)


def safe_name(s: str) -> str:
    for ch in r'\/:*?"<>|':
        s = s.replace(ch, "_")
    return s.strip().rstrip(".")[:60].strip()


# ---------------------------------------------------------------------------
# XML parser
# ---------------------------------------------------------------------------

def parse_movimento_economico_xml(xml_bytes: bytes) -> list[dict]:
    try:
        r = ET.fromstring(xml_bytes)
    except ET.ParseError:
        return []

    if r.tag != "movimento_economico":
        return []

    toma      = r.find("tomador")
    toma_cnpj = (toma.findtext("cpf_cnpj") or "").strip() if toma is not None else ""
    toma_nome = (toma.findtext("razao_social") or "").strip() if toma is not None else ""
    toma_cmc  = (toma.findtext("cmc") or "").strip() if toma is not None else ""

    results = []
    for nota in r.findall(".//nota"):
        prest      = nota.find("prestador")
        prest_cnpj = (prest.findtext("cpf_cnpj") or "").strip() if prest is not None else ""
        prest_nome = (prest.findtext("razao_social") or "").strip() if prest is not None else ""

        v_irrf   = fm(nota.findtext("valor_ir"))
        v_pis    = fm(nota.findtext("valor_pis"))
        v_cofins = fm(nota.findtext("valor_cofins"))
        v_inss   = fm(nota.findtext("valor_inss"))
        v_csll   = fm(nota.findtext("valor_csll"))
        v_iss    = fm(nota.findtext("valor_iss"))
        v_serv   = fm(nota.findtext("valor_nota"))
        v_bc     = fm(nota.findtext("base_calculo"))
        v_liq    = fm(nota.findtext("valor_liquido"))

        tipo_nota  = nota.findtext("tipo_nota", "")
        tp_ret_iss = "2" if tipo_nota == "TOMADOR" and v_iss > 0 else "1"
        v_iss_ret  = v_iss if tp_ret_iss == "2" else 0.0
        v_total_ret = v_iss_ret + v_irrf + v_pis + v_cofins + v_inss + v_csll

        situacao = nota.findtext("situacao", "")
        if situacao and situacao.upper() in ("C", "CANCELADA", "CANCELADO"):
            continue

        n_nfse   = nota.get("numero") or nota.findtext("rps", "")
        rps      = nota.findtext("rps", "")
        d_compet = nota.findtext("data_emissao", "")
        xloc     = nota.findtext("cidade_prestacao", "")
        xdesc    = (nota.findtext("nome_atividade") or "")[:200]

        results.append({
            "toma_cnpj":    toma_cnpj,
            "toma_nome":    toma_nome,
            "toma_cmc":     toma_cmc,
            "prest_cnpj":   prest_cnpj,
            "prest_nome":   prest_nome,
            "n_nfse":       n_nfse,
            "rps":          rps,
            "d_compet":     d_compet,
            "xloc":         xloc,
            "xdesc":        xdesc,
            "v_serv":       v_serv,
            "v_bc":         v_bc,
            "v_issqn":      v_iss,
            "tp_ret_iss":   tp_ret_iss,
            "v_iss_ret":    v_iss_ret,
            "v_irrf":       v_irrf,
            "v_csll":       v_csll,
            "v_pis":        v_pis,
            "v_cofins":     v_cofins,
            "v_inss":       v_inss,
            "v_total_ret":  v_total_ret,
            "v_liq":        v_liq,
            "xml_element":  ET.tostring(nota, encoding="unicode"),
        })

    return results


# ---------------------------------------------------------------------------
# Output: folder structure + XLSX
# ---------------------------------------------------------------------------

def write_company_xlsx(company_name: str, notas: list[dict], out_dir: Path) -> Path:
    wb   = Workbook()
    hf   = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    tf   = Font(name="Arial", bold=True, size=11, color="1A56A0")
    nf   = Font(name="Arial", size=10)
    bf   = Font(name="Arial", bold=True, size=10)
    af   = PatternFill("solid", start_color="F0F5FC")
    wf   = PatternFill("solid", start_color="FFFFFF")
    totf = PatternFill("solid", start_color="D6E4F7")
    ctr  = Alignment(horizontal="center", vertical="center")
    lft  = Alignment(horizontal="left",   vertical="center")
    rgt  = Alignment(horizontal="right",  vertical="center")
    thin = Side(style="thin", color="BFBFBF")
    bdr  = Border(left=thin, right=thin, top=thin, bottom=thin)
    money = "#,##0.00"

    HEADERS = ["Nº NFSe","Data","Prestador","CNPJ Prestador",
               "Vl. Serviço","ISS Ret.","PIS","COFINS",
               "IR","CSLL","INSS","Total Ret.","Descrição","Cidade"]
    WIDTHS  = {"A":10,"B":12,"C":40,"D":20,"E":14,"F":12,"G":12,
               "H":14,"I":12,"J":12,"K":12,"L":14,"M":50,"N":20}
    MONEY_COLS = {5,6,7,8,9,10,11,12}
    CENTER_COLS = {1,2,4}

    def make_sheet(ws, title_text, rows, hfill_color):
        hfill = PatternFill("solid", start_color=hfill_color)
        ws.merge_cells("A1:N1")
        ws["A1"] = title_text
        ws["A1"].font = tf; ws["A1"].alignment = lft
        ws.row_dimensions[1].height = 20

        for col, h in enumerate(HEADERS, 1):
            c = ws.cell(row=2, column=col, value=h)
            c.font = hf; c.fill = hfill; c.alignment = ctr; c.border = bdr
        ws.row_dimensions[2].height = 16

        for i, n in enumerate(rows, start=3):
            fill = af if i % 2 == 0 else wf
            vals = [n["n_nfse"], n["d_compet"], n["prest_nome"], n["prest_cnpj"],
                    n["v_serv"], n["v_iss_ret"], n["v_pis"], n["v_cofins"],
                    n["v_irrf"], n["v_csll"], n["v_inss"], n["v_total_ret"],
                    n["xdesc"], n["xloc"]]
            for col, val in enumerate(vals, 1):
                c = ws.cell(row=i, column=col, value=val)
                c.font = nf; c.fill = fill; c.border = bdr
                if col in MONEY_COLS:
                    c.number_format = money; c.alignment = rgt
                elif col in CENTER_COLS:
                    c.alignment = ctr
                else:
                    c.alignment = lft

        tr = len(rows) + 3
        ws.cell(row=tr, column=1, value="TOTAIS").font = bf
        ws.cell(row=tr, column=1).fill = totf
        ws.cell(row=tr, column=1).border = bdr
        ws.cell(row=tr, column=1).alignment = lft
        for col, key in [(5,"v_serv"),(6,"v_iss_ret"),(7,"v_pis"),(8,"v_cofins"),
                          (9,"v_irrf"),(10,"v_csll"),(11,"v_inss"),(12,"v_total_ret")]:
            c = ws.cell(row=tr, column=col, value=sum(n[key] for n in rows))
            c.font = bf; c.fill = totf; c.number_format = money
            c.alignment = rgt; c.border = bdr
        for col in (2,3,4,13,14):
            c = ws.cell(row=tr, column=col); c.fill = totf; c.border = bdr

        for l, w in WIDTHS.items():
            ws.column_dimensions[l].width = w
        ws.freeze_panes = "A3"

    ws1 = wb.active
    ws1.title = "Todas as Notas"
    make_sheet(ws1, f"TODAS AS NOTAS — {company_name}", notas, "1A56A0")

    ret_notas = [n for n in notas if has_federal_retencao(n)]
    ws2 = wb.create_sheet("Retenção Federal")
    make_sheet(ws2, f"RETENÇÃO FEDERAL — {company_name}", ret_notas, "C0392B")

    out_path = out_dir / f"NFS-e_{safe_name(company_name)}.xlsx"
    wb.save(str(out_path))
    return out_path


def save_xml_files(notas: list[dict], xml_dir: Path, ret_dir: Path):
    xml_dir.mkdir(parents=True, exist_ok=True)
    ret_dir.mkdir(parents=True, exist_ok=True)
    seen = {}
    for n in notas:
        num   = n["n_nfse"] or n["rps"] or "SN"
        fname = f"NFSe_{num}.xml"
        key   = fname
        if key in seen and seen[key] != n["prest_cnpj"]:
            fname = f"NFSe_{num}_{re.sub(r'[^0-9]','',n['prest_cnpj'])[-8:]}.xml"
        seen[key] = n["prest_cnpj"]
        content = (f'<?xml version="1.0" encoding="UTF-8"?>\n'
                   f'<nota_nfse n_nfse="{n["n_nfse"]}">\n{n["xml_element"]}\n</nota_nfse>')
        (xml_dir / fname).write_text(content, encoding="utf-8")
        if has_federal_retencao(n):
            (ret_dir / fname).write_text(content, encoding="utf-8")


# ---------------------------------------------------------------------------
# Portal helpers (reuse login/company list from iss_worker)
# ---------------------------------------------------------------------------

def login(user, password, visible=False):
    from iss_worker import login as _login
    return _login(user, password, visible)


def get_company_list(page):
    from iss_worker import get_company_list as _get
    return _get(page)


def go_to_empresas(page):
    from iss_worker import go_to_empresas as _go
    _go(page)


def download_nfse_xml(page, company_id: str, portal_period: str) -> bytes | None:
    """
    Navigate to company, go to Períodos > Exportar XML, POST gerar_xml.php,
    return raw XML bytes or None if no data.

    portal_period: YYYYM format (e.g. '20266' for Jun 2026)
    """
    go_to_empresas(page)
    page.wait_for_selector("select#clientes", timeout=10000)
    page.select_option("select#clientes", company_id)
    page.wait_for_timeout(800)

    page.click("input[value='Acessar Movimento Econômico']")
    page.wait_for_load_state("networkidle", timeout=30000)
    page.wait_for_timeout(1000)

    # Navigate to Períodos > Exportar XML
    page.evaluate("() => { mostrarConteudo('iss-exportar_dados.php'); }")
    page.wait_for_timeout(1500)

    # Check the period exists in the dropdown
    try:
        page.wait_for_selector("select#anomes_ini", timeout=8000)
        options = page.eval_on_selector_all(
            "select#anomes_ini option",
            "opts => opts.map(o => o.value)"
        )
        if portal_period not in options:
            return None
    except Exception:
        return None

    # Use page.request (Playwright's fetch API with session cookies) to POST
    # This avoids dealing with file download interception
    try:
        response = page.request.post(
            f"{BASE}/gerar_xml.php",
            form={
                "anomes_ini":    portal_period,
                "anomes_fim":    portal_period,
                "lancamentos":   "T",
                "ordem":         "DESC",
                "submit":        "Gerar Arquivo",
                "form_abertura": "1",
            }
        )
        if response.status != 200:
            return None
        body = response.body()
        # Empty response = no data for this period
        if not body or len(body) < 50:
            return None
        return body
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Main generator
# ---------------------------------------------------------------------------

def run_nfse_pm(user: str, password: str, target_period: str,
                selected_ids: list[str] | None, output_base: str,
                visible: bool = False):
    """
    Generator — yields log strings.
    Ends with __DONE__:code and __SUMMARY__:json.
    target_period: MMYYYY (UI format), converted internally.
    selected_ids: list of CMC ids, None = all companies.
    """
    import json

    def log(msg):
        return f"[{datetime.now().strftime('%H:%M:%S')}] {msg}"

    portal_period = to_portal_period(target_period)
    period_label  = f"{target_period[:2]}/{target_period[2:]}"

    yield log("Iniciando browser...")
    try:
        p, browser, context, page = login(user, password, visible)
    except Exception as e:
        yield log(f"ERRO — Login: {e}")
        yield "__DONE__:1"
        return

    yield log("Login efetuado. Carregando empresas...")
    try:
        companies = get_company_list(page)
    except Exception as e:
        yield log(f"ERRO — Lista de empresas: {e}")
        browser.close(); p.stop()
        yield "__DONE__:1"
        return

    if selected_ids:
        companies = [c for c in companies if c["id"] in selected_ids]

    yield log(f"{len(companies)} empresas | Período: {period_label} (portal: {portal_period})")

    out_base = Path(output_base) / f"{target_period[:2]}-{target_period[2:]}"
    out_base.mkdir(parents=True, exist_ok=True)

    total   = len(companies)
    stats   = {"ok": 0, "no_data": 0, "error": 0, "federal": 0}
    summary = []

    try:
        for i, comp in enumerate(companies, 1):
            cid  = comp["id"]
            name = comp["name"]
            yield log(f"[{i}/{total}] {name}")

            try:
                xml_bytes = download_nfse_xml(page, cid, portal_period)

                if not xml_bytes:
                    yield log(f"  -> Sem dados para {period_label}")
                    stats["no_data"] += 1
                    continue

                notas = parse_movimento_economico_xml(xml_bytes)
                if not notas:
                    yield log(f"  -> XML vazio ou sem notas válidas")
                    stats["no_data"] += 1
                    continue

                fed = [n for n in notas if has_federal_retencao(n)]
                fed_str = f" | {len(fed)} com retenção federal ✓" if fed else ""
                yield log(f"  -> {len(notas)} notas{fed_str}")

                company_dir = out_base / safe_name(name)
                company_dir.mkdir(parents=True, exist_ok=True)

                # Save raw XML
                (company_dir / f"mov_economico_{cid}_{portal_period}.xml").write_bytes(xml_bytes)

                # Save individual nota XMLs
                save_xml_files(notas, company_dir / "Todas as Notas", company_dir / "Retenção Federal")

                # Write XLSX
                write_company_xlsx(name, notas, company_dir)

                stats["ok"] += 1
                stats["federal"] += len(fed)
                summary.append({"name": name, "cmc": cid, "total": len(notas), "federal": len(fed)})

            except Exception as e:
                yield log(f"  -> ERRO: {e}")
                stats["error"] += 1

            if i < total:
                time.sleep(random.uniform(0.8, 2.0))

    except GeneratorExit:
        pass
    finally:
        browser.close()
        p.stop()

    yield log(f"Concluído — {stats['ok']} OK | {stats['federal']} notas federais | {stats['no_data']} sem dados | {stats['error']} erros")
    yield "__DONE__:0"
    yield f"__SUMMARY__:{json.dumps(summary, ensure_ascii=False)}"
